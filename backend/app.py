from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from docx import Document
from datetime import datetime
from docx2pdf import convert
import os
import uuid

app = Flask(__name__)
CORS(app)  # Allow requests from frontend

@app.route('/api/generate', methods=['POST'])
def generate_certificate():
    data = request.json
    name = data.get('name')
    domain = data.get('domain')
    start_date = data.get('start_date')
    end_date = data.get('end_date')
    gender = data.get('gender', 'other').lower()

    # Pronoun mapping
    pronouns = {"male": ("he", "him"), "female": ("she", "her"), "other": ("they", "them")}
    he_she, him_her = pronouns.get(gender, ("they", "them"))

    issued_date = datetime.today().strftime('%B %d, %Y')

    # Generate unique filenames to avoid conflicts
    uid = str(uuid.uuid4())
    docx_filename = f"certificate_{uid}.docx"
    pdf_filename = f"certificate_{uid}.pdf"

    # Load and replace placeholders
    doc = Document("SpectoV_Cert.docx")
    for para in doc.paragraphs:
        para.text = para.text.replace("{{Domain}}", domain)
        para.text = para.text.replace("{{Start Date}}", start_date)
        para.text = para.text.replace("{{End Date}}", end_date)
        para.text = para.text.replace("{{he/she/they}}", he_she)
        para.text = para.text.replace("{{him/her/them}}", him_her)
        para.text = para.text.replace("{{Name}}", name)
        para.text = para.text.replace("ISSUED DATE :", f"ISSUED DATE : {issued_date}")

    doc.save(docx_filename)
    convert(docx_filename, pdf_filename)

    # Send file and cleanup
    response = send_file(pdf_filename, as_attachment=True)

    @response.call_on_close
    def cleanup():
        os.remove(docx_filename)
        os.remove(pdf_filename)

    return response

if __name__ == '__main__':
    app.run(debug=True, port=5001)
